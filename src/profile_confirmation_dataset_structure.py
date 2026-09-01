from __future__ import annotations

import hashlib
import json
import time
from collections import Counter
from datetime import date, datetime, time as datetime_time
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl import load_workbook

from public_paths import repo_path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = PROJECT_ROOT / "data" / "raw" / "confirmation_online_retail_ii" / "online_retail_II.xlsx"
ACQUISITION_PATH = PROJECT_ROOT / "reports" / "bizhallu_confirmation_dataset_acquisition_report.json"
REPORT_PATH = PROJECT_ROOT / "reports" / "bizhallu_confirmation_dataset_structure_report.json"

WINDOW_START = datetime(2009, 12, 1)
WINDOW_END = datetime(2010, 12, 1)
DOCUMENTED_INSTANCE_COUNT = 1_067_371
TYPE_SAMPLE_LIMIT_PER_COLUMN_PER_SHEET = 2_000

CANONICAL_COLUMNS = [
    "InvoiceNo",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "UnitPrice",
    "CustomerID",
    "Country",
]

HEADER_ALIASES = {
    "Invoice": "InvoiceNo",
    "InvoiceNo": "InvoiceNo",
    "StockCode": "StockCode",
    "Description": "Description",
    "Quantity": "Quantity",
    "InvoiceDate": "InvoiceDate",
    "Price": "UnitPrice",
    "UnitPrice": "UnitPrice",
    "Customer ID": "CustomerID",
    "CustomerID": "CustomerID",
    "Country": "Country",
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def iso(value: datetime | None) -> str | None:
    return value.isoformat(sep=" ") if value else None


def value_type(value: Any) -> str:
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, str):
        return "str"
    return type(value).__name__


def parse_source_datetime(value: Any) -> tuple[datetime | None, str | None]:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None), None
    if isinstance(value, date):
        return datetime.combine(value, datetime_time.min), None
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return None, "blank_string"
        try:
            return datetime.fromisoformat(stripped), None
        except ValueError:
            return None, "unparseable_string"
    if value is None:
        return None, "null"
    return None, f"unsupported_type:{type(value).__name__}"


def main() -> None:
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"Missing local workbook: {XLSX_PATH}")
    if not ACQUISITION_PATH.exists():
        raise FileNotFoundError(f"Missing acquisition report: {ACQUISITION_PATH}")

    acquisition = json.loads(ACQUISITION_PATH.read_text(encoding="utf-8"))
    workbook_sha256 = sha256_file(XLSX_PATH)
    expected_sha256 = acquisition["raw_artifacts"]["xlsx"]["sha256"]
    if workbook_sha256 != expected_sha256:
        raise RuntimeError("Workbook SHA-256 does not match the acquisition report.")

    started = time.perf_counter()
    workbook = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    sheet_profiles: list[dict[str, Any]] = []
    all_actual_headers: list[list[str]] = []
    warnings: list[dict[str, Any]] = []

    total_data_rows = 0
    total_fully_blank_rows = 0
    total_valid_date_rows = 0
    total_blank_date_rows = 0
    total_invalid_date_rows = 0
    total_before_start = 0
    total_in_window = 0
    total_at_or_after_end = 0
    overall_date_min: datetime | None = None
    overall_date_max: datetime | None = None
    strict_date_min: datetime | None = None
    strict_date_max: datetime | None = None

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if raw_headers is None:
            raise RuntimeError(f"Worksheet {sheet_name!r} is empty.")
        actual_headers = [str(value).strip() if value is not None else "" for value in raw_headers]
        all_actual_headers.append(actual_headers)
        canonical_headers = [HEADER_ALIASES.get(value) for value in actual_headers]
        if any(value is None for value in canonical_headers):
            raise RuntimeError(
                f"Worksheet {sheet_name!r} has unmapped headers: "
                f"{[actual_headers[index] for index, value in enumerate(canonical_headers) if value is None]!r}."
            )
        if canonical_headers != CANONICAL_COLUMNS:
            raise RuntimeError(
                f"Worksheet {sheet_name!r} canonical header order is {canonical_headers!r}, "
                f"expected {CANONICAL_COLUMNS!r}."
            )
        if len(set(actual_headers)) != len(actual_headers):
            raise RuntimeError(f"Worksheet {sheet_name!r} contains duplicate headers.")

        date_index = canonical_headers.index("InvoiceDate")
        sample_type_counts = {column: Counter() for column in CANONICAL_COLUMNS}
        sampled_non_null_counts = {column: 0 for column in CANONICAL_COLUMNS}

        sheet_data_rows = 0
        sheet_fully_blank_rows = 0
        sheet_valid_dates = 0
        sheet_blank_dates = 0
        sheet_invalid_dates = 0
        sheet_before_start = 0
        sheet_in_window = 0
        sheet_at_or_after_end = 0
        sheet_date_min: datetime | None = None
        sheet_date_max: datetime | None = None
        sheet_strict_min: datetime | None = None
        sheet_strict_max: datetime | None = None

        for row_number, row in enumerate(rows, start=2):
            if not row or all(value is None for value in row):
                sheet_fully_blank_rows += 1
                continue
            sheet_data_rows += 1

            for index, column in enumerate(CANONICAL_COLUMNS):
                value = row[index] if index < len(row) else None
                if value is not None and sampled_non_null_counts[column] < TYPE_SAMPLE_LIMIT_PER_COLUMN_PER_SHEET:
                    sample_type_counts[column][value_type(value)] += 1
                    sampled_non_null_counts[column] += 1

            raw_date = row[date_index] if date_index < len(row) else None
            parsed_date, date_error = parse_source_datetime(raw_date)
            if parsed_date is None:
                if date_error in {"null", "blank_string"}:
                    sheet_blank_dates += 1
                else:
                    sheet_invalid_dates += 1
                if len(warnings) < 20:
                    warnings.append(
                        {
                            "sheet": sheet_name,
                            "row_number": row_number,
                            "field": "InvoiceDate",
                            "warning": date_error,
                        }
                    )
                continue

            sheet_valid_dates += 1
            sheet_date_min = parsed_date if sheet_date_min is None else min(sheet_date_min, parsed_date)
            sheet_date_max = parsed_date if sheet_date_max is None else max(sheet_date_max, parsed_date)
            if parsed_date < WINDOW_START:
                sheet_before_start += 1
            elif parsed_date >= WINDOW_END:
                sheet_at_or_after_end += 1
            else:
                sheet_in_window += 1
                sheet_strict_min = parsed_date if sheet_strict_min is None else min(sheet_strict_min, parsed_date)
                sheet_strict_max = parsed_date if sheet_strict_max is None else max(sheet_strict_max, parsed_date)

        total_data_rows += sheet_data_rows
        total_fully_blank_rows += sheet_fully_blank_rows
        total_valid_date_rows += sheet_valid_dates
        total_blank_date_rows += sheet_blank_dates
        total_invalid_date_rows += sheet_invalid_dates
        total_before_start += sheet_before_start
        total_in_window += sheet_in_window
        total_at_or_after_end += sheet_at_or_after_end
        if sheet_date_min is not None:
            overall_date_min = sheet_date_min if overall_date_min is None else min(overall_date_min, sheet_date_min)
        if sheet_date_max is not None:
            overall_date_max = sheet_date_max if overall_date_max is None else max(overall_date_max, sheet_date_max)
        if sheet_strict_min is not None:
            strict_date_min = sheet_strict_min if strict_date_min is None else min(strict_date_min, sheet_strict_min)
        if sheet_strict_max is not None:
            strict_date_max = sheet_strict_max if strict_date_max is None else max(strict_date_max, sheet_strict_max)

        sheet_profiles.append(
            {
                "sheet_name": sheet_name,
                "declared_max_row_including_header": sheet.max_row,
                "declared_max_column": sheet.max_column,
                "actual_headers": actual_headers,
                "canonical_headers": canonical_headers,
                "data_rows": sheet_data_rows,
                "fully_blank_rows": sheet_fully_blank_rows,
                "valid_invoice_date_rows": sheet_valid_dates,
                "blank_invoice_date_rows": sheet_blank_dates,
                "invalid_invoice_date_rows": sheet_invalid_dates,
                "date_min": iso(sheet_date_min),
                "date_max": iso(sheet_date_max),
                "strict_window_row_count": sheet_in_window,
                "strict_window_date_min": iso(sheet_strict_min),
                "strict_window_date_max": iso(sheet_strict_max),
                "rows_before_window": sheet_before_start,
                "rows_at_or_after_window": sheet_at_or_after_end,
                "sampled_non_null_value_type_counts": {
                    column: dict(sorted(sample_type_counts[column].items())) for column in CANONICAL_COLUMNS
                },
                "type_sample_limit_per_column": TYPE_SAMPLE_LIMIT_PER_COLUMN_PER_SHEET,
            }
        )

    workbook.close()
    elapsed = round(time.perf_counter() - started, 3)
    unique_header_sets = {tuple(headers) for headers in all_actual_headers}
    actual_headers = list(next(iter(unique_header_sets))) if len(unique_header_sets) == 1 else []
    exact_metadata_names = actual_headers == CANONICAL_COLUMNS

    report = {
        "status": "workbook_structure_and_strict_window_verified",
        "study_name": "Confirmation Set v1",
        "candidate_id": "uci_online_retail_ii_prior_period",
        "profile_scope": "workbook structure, sampled value types, and complete InvoiceDate boundary scan only",
        "acquisition_report_path": repo_path(ACQUISITION_PATH),
        "workbook_path": repo_path(XLSX_PATH),
        "workbook_sha256": workbook_sha256,
        "parser": {"library": "openpyxl", "version": openpyxl.__version__, "read_only": True, "data_only": True},
        "documented_instance_count": DOCUMENTED_INSTANCE_COUNT,
        "workbook_sheet_count": len(sheet_profiles),
        "sheet_names": [item["sheet_name"] for item in sheet_profiles],
        "workbook_total_data_rows": total_data_rows,
        "documented_instance_delta": total_data_rows - DOCUMENTED_INSTANCE_COUNT,
        "workbook_total_fully_blank_rows": total_fully_blank_rows,
        "canonical_columns": CANONICAL_COLUMNS,
        "actual_workbook_headers": actual_headers,
        "header_alias_mapping": {header: HEADER_ALIASES[header] for header in actual_headers},
        "metadata_header_drift_detected": not exact_metadata_names,
        "metadata_header_drift": {
            "severity": "low_documented_alias_mapping_required",
            "actual_to_canonical_changes": {
                header: HEADER_ALIASES[header]
                for header in actual_headers
                if header != HEADER_ALIASES[header]
            },
            "impact": "Existing Online Retail scripts cannot assume identical raw header names; ingestion must rename aliases before applying shared business rules.",
        },
        "date_policy": {
            "source_timezone": "not documented; workbook datetimes are treated as naive source-local timestamps",
            "start_inclusive": WINDOW_START.isoformat(),
            "end_exclusive": WINDOW_END.isoformat(),
        },
        "date_scan": {
            "valid_invoice_date_rows": total_valid_date_rows,
            "blank_invoice_date_rows": total_blank_date_rows,
            "invalid_invoice_date_rows": total_invalid_date_rows,
            "overall_date_min": iso(overall_date_min),
            "overall_date_max": iso(overall_date_max),
            "rows_before_window": total_before_start,
            "strict_window_row_count": total_in_window,
            "strict_window_date_min": iso(strict_date_min),
            "strict_window_date_max": iso(strict_date_max),
            "rows_at_or_after_window": total_at_or_after_end,
        },
        "sheet_profiles": sheet_profiles,
        "parse_warning_count": total_blank_date_rows + total_invalid_date_rows,
        "parse_warning_examples": warnings,
        "full_missingness_profile_complete": False,
        "duplicate_and_grain_profile_complete": False,
        "business_rule_profile_complete": False,
        "historical_overlap_check_complete": False,
        "context_feasibility_check_complete": False,
        "context_manifest_created": False,
        "prompt_created": False,
        "model_run_performed": False,
        "new_metrics_reported": False,
        "execution_ready": False,
        "no_new_results": True,
        "next_authorized_action": "Run the full field-level quality, duplicate/grain, cancellation, overlap, and context-feasibility profile without creating prompts.",
    }
    REPORT_PATH.write_text(json.dumps(report, indent=2, ensure_ascii=True), encoding="utf-8")
    print(
        json.dumps(
            {
                "status": report["status"],
                "workbook_total_data_rows": total_data_rows,
                "sheet_names": report["sheet_names"],
                "metadata_header_drift_detected": report["metadata_header_drift_detected"],
                "strict_window_row_count": total_in_window,
                "overall_date_min": report["date_scan"]["overall_date_min"],
                "overall_date_max": report["date_scan"]["overall_date_max"],
                "parse_warning_count": report["parse_warning_count"],
                "scan_elapsed_seconds": elapsed,
                "report_path": repo_path(REPORT_PATH),
            },
            indent=2,
            ensure_ascii=True,
        )
    )


if __name__ == "__main__":
    main()
